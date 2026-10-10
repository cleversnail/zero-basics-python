"""第 94 讲 · 理论盘点：文件类型与对应的 Python 函数库

本讲把「什么文件格式，用什么 Python 库」讲清楚，并真跑一遍：
    1. xlsx（新版 Excel）  -> 用 openpyxl 读写
    2. docx（新版 Word）   -> 用 python-docx 读写（导入名是 import docx）
    3. 顺带完成本讲课后作业：把 Excel 第一张表格的内容写进一个 docx 文件

需要先安装依赖（本机已安装，能真跑）：
    pip install openpyxl python-docx

运行：python3 094-file-types.py
"""

from openpyxl import Workbook, load_workbook
from docx import Document

print("=" * 48)
print("1) 用 openpyxl 新建并读回一个 xlsx")
print("=" * 48)

# Workbook() 新建一个空的 Excel 工作簿；wb.active 拿到当前活动工作表
wb = Workbook()
ws = wb.active
ws.title = "订单"

# 一行一行追加数据（列表的每个元素对应一格）
ws.append(["订单号", "客户", "金额"])
ws.append([1001, "张三", 250])
ws.append([1002, "李四", 180])

# save() 才真正落盘，生成 094-demo.xlsx 文件
wb.save("094-demo.xlsx")

# load_workbook() 再把文件读回来，验证刚才真的写进去了
wb2 = load_workbook("094-demo.xlsx")
print("工作表名字：", wb2.sheetnames)
for row in wb2.active.iter_rows(values_only=True):
    print("  ", row)

print()
print("=" * 48)
print("2) 用 python-docx 新建并读回一个 docx")
print("=" * 48)

# Document() 新建一个空的 Word 文档
doc = Document()
doc.add_heading("第 94 讲演示文档", level=1)   # level=1 是一级标题
doc.add_paragraph("这一行是普通段落。")
doc.save("094-demo.docx")                     # 落盘

# 读回来：document.paragraphs 是所有段落
doc2 = Document("094-demo.docx")
for p in doc2.paragraphs:
    if p.text.strip():
        print("  段落：", p.text)

print()
print("=" * 48)
print("3) 课后作业：把 Excel 第一张表写进 docx")
print("=" * 48)

# 思路：先用 openpyxl 读 Excel 的表格 -> 再用 python-docx 写成 Word 表格
src = load_workbook("094-demo.xlsx")
rows = list(src.active.iter_rows(values_only=True))
print("从 Excel 读到的数据：")
for r in rows:
    print("  ", r)

out = Document()
out.add_heading("订单汇总", level=1)

# add_table(rows=1, cols=列数) 先建一个只有表头的表
table = out.add_table(rows=1, cols=len(rows[0]))
hdr = table.rows[0].cells
for i, name in enumerate(rows[0]):
    hdr[i].text = str(name)

# 剩下的每一行：add_row() 新增一行，再逐格填文本
for data_row in rows[1:]:
    cells = table.add_row().cells
    for i, value in enumerate(data_row):
        cells[i].text = str(value)

out.save("094-作业-excel转docx.docx")
print("已保存 094-作业-excel转docx.docx（用 Word 打开即可看到这张表）")

# 读回作业文件，确认表格真的写进去了
check = Document("094-作业-excel转docx.docx")
print("回读 094-作业-excel转docx.docx 的表格内容：")
for t in check.tables:
    for row in t.rows:
        print("  ", [c.text for c in row.cells])

print()
print("说明：xls（旧版 Excel）不能用 openpyxl，读用 xlrd、写用 xlwt——本机未装，故只讲不跑。")
