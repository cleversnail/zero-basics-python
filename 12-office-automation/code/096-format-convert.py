"""第 96 讲 · 如何实现数据的批量格式转换？

本讲真跑两种常见的格式转换：
    1. Word -> Excel：读 sample.docx 里的文字与表格，写进一个 xlsx
    2. Excel -> Word：读一个 xlsx，写进一个 docx 表格
    （第 1 种正是本讲课后作业：把指定 Word 文件的内容转成同名 Excel）

需要先安装依赖（本机已安装，能真跑）：pip install openpyxl python-docx

运行：python3 096-format-convert.py
"""

from openpyxl import Workbook, load_workbook
from docx import Document

print("=" * 48)
print("1) Word -> Excel：把 sample.docx 的文字和表格转成 xlsx")
print("=" * 48)

# 读 Word
doc = Document("sample.docx")

# 收集非空段落文本（正文）
paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
print("Word 里的段落：")
for t in paragraphs:
    print("  ", t)

# 收集表格内容
table_rows = []
for t in doc.tables:
    for row in t.rows:
        table_rows.append([c.text for c in row.cells])
print("Word 里的表格：")
for r in table_rows:
    print("  ", r)

# 写 Excel：正文一列，表格另起一张工作表
wb = Workbook()
ws1 = wb.active
ws1.title = "正文"
for t in paragraphs:
    ws1.append([t])

ws2 = wb.create_sheet("表格")
for r in table_rows:
    ws2.append(r)

wb.save("096-word转excel.xlsx")
print("已生成 096-word转excel.xlsx，工作表：", wb.sheetnames)

# 读回验证
back = load_workbook("096-word转excel.xlsx")
print("回读「正文」表：")
for row in back["正文"].iter_rows(values_only=True):
    print("  ", row)
print("回读「表格」表：")
for row in back["表格"].iter_rows(values_only=True):
    print("  ", row)

print()
print("=" * 48)
print("2) Excel -> Word：把刚才的 xlsx 表格写回一个 docx")
print("=" * 48)

# 用 openpyxl 读表格，用 python-docx 写表格（转换关键：确定好列数/列宽）
src = load_workbook("096-word转excel.xlsx")["表格"]
rows = list(src.iter_rows(values_only=True))

out = Document()
out.add_heading("由 Excel 转换而来", level=1)

cols = len(rows[0])
table = out.add_table(rows=1, cols=cols)
for i, name in enumerate(rows[0]):
    table.rows[0].cells[i].text = str(name)
for data_row in rows[1:]:
    cells = table.add_row().cells
    for i, value in enumerate(data_row):
        cells[i].text = str(value)

out.save("096-excel转word.docx")
print("已生成 096-excel转word.docx")

# 读回验证
check = Document("096-excel转word.docx")
for t in check.tables:
    for row in t.rows:
        print("  ", [c.text for c in row.cells])

print()
print("提示：格式转换常要同时开多个文件，操作完记得 save() 落盘；")
print("      openpyxl 的 workbook 对象也要及时让变量失效，否则内存占用会升高。")
